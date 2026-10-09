"""
Export report: filters (including the joining-date one that used to match nothing), the Excel
workbook's real contents, and company scoping for HR Managers.
"""
import io
import re
import uuid
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.models import (
    UserRole, RequestStatus, OnboardingRequest, ApprovalAction, ApprovalActionType, AuditLog,
    FormField, FieldType,
)
from .conftest import login, _make_user


@pytest.fixture(autouse=True)
def _app_ctx(app):
    with app.app_context():
        yield


def _req(db, ini_id, name, status=RequestStatus.ACTIVE, company="RDC", plant="Plant A", desig="Tech",
         form=None, retry=0, created=None):
    r = OnboardingRequest(initiated_by=ini_id, status=status, public_token=uuid.uuid4().hex, candidate_name=name,
                          company_code=company, plant_location=plant, designation=desig, retry_count=retry)
    if created:
        r.created_at = created
    db.session.add(r)
    db.session.flush()
    r.form_data = {"associate_name": name, **(form or {})}
    db.session.flush()
    return r


def _preview(client, qs=""):
    r = client.get("/exports/active-employees/preview" + qs)
    assert r.status_code == 200
    return r.get_json()


def _cands(data):
    cols = data["columns"]
    i = cols.index("Associate Name (as per Aadhaar)")
    return sorted(row[i] for row in data["rows"])


@pytest.fixture
def world(db):
    admin = _make_user("Exp Admin", "expadmin@t.com", UserRole.SUPER_ADMIN, db)
    ini = _make_user("Exp Ini", "expini@t.com", UserRole.INITIATOR, db, companies=["RDC", "ROBO"])
    # the report's form columns come from the configured FormFields (the test DB isn't seeded)
    for i, (key, label) in enumerate((("associate_name", "Associate Name (as per Aadhaar)"),
                                       ("contract_from", "Contract From (Date of Joining)"))):
        db.session.add(FormField(field_key=key, field_label=label, field_type=FieldType.TEXT, step=1, sort_order=i))
    db.session.commit()
    return admin.id, ini.id


class TestExportFilters:
    def test_default_is_approved_only(self, client, db, world):
        aid, ini = world
        _req(db, ini, "Approved A")
        _req(db, ini, "Pending B", status=RequestStatus.PENDING_BH)
        db.session.commit()
        login(client, "expadmin@t.com")
        data = _preview(client)
        assert data["total"] == 1 and data["title"] == "Approved — Onboarding Report"
        assert data["chips"][0] == {"label": "Status", "value": "Approved (default)"}

    def test_status_multi_select_and_labels(self, client, db, world):
        aid, ini = world
        _req(db, ini, "S1", status=RequestStatus.PENDING_BH)
        _req(db, ini, "S2", status=RequestStatus.REJECTED_BH)
        _req(db, ini, "S3", status=RequestStatus.ACTIVE)
        db.session.commit()
        login(client, "expadmin@t.com")
        data = _preview(client, "?status=PENDING_BH&status=REJECTED_BH")
        assert data["total"] == 2
        assert data["title"] == "Pending Business / Functional Head + Rejected by Business / Functional Head — Onboarding Report"
        assert _preview(client, "?status=BOGUS")["total"] in (0, 1, 3)       # junk never breaks the report

    def test_joining_date_filter_uses_the_real_form_field(self, client, db, world):
        aid, ini = world
        _req(db, ini, "Join Early", form={"contract_from": "2026-01-10"})
        _req(db, ini, "Join Mid", form={"contract_from": "2026-06-15"})
        _req(db, ini, "Join Late", form={"contract_from": "2026-12-01"})
        _req(db, ini, "Join None")
        db.session.commit()
        login(client, "expadmin@t.com")
        data = _preview(client, "?joining_from=2026-06-01&joining_to=2026-06-30")
        assert data["total"] == 1 and _cands(data) == ["Join Mid"]
        assert _preview(client, "?joining_from=2026-06-01")["total"] == 2
        assert _preview(client, "?joining_to=2026-02-01")["total"] == 1
        chips = _preview(client, "?joining_from=2026-06-01&joining_to=2026-06-30")["chips"]
        assert {"label": "Joining date", "value": "01/06/2026 → 30/06/2026"} in chips
        assert _preview(client, "?joining_from=not-a-date")["total"] == 4

    def test_plant_designation_company_retry_and_submission_date(self, client, db, world):
        aid, ini = world
        _req(db, ini, "F1", plant="Plant X", desig="Welder", company="RDC", retry=0)
        _req(db, ini, "F2", plant="Plant Y", desig="Fitter", company="ROBO", retry=3)
        _req(db, ini, "F3", plant="Plant Y", desig="Fitter", company="ROBO", retry=1,
             created=datetime.utcnow() - timedelta(days=40))
        db.session.commit()
        login(client, "expadmin@t.com")
        assert _cands(_preview(client, "?plant=Plant X")) == ["F1"]
        assert _cands(_preview(client, "?designation=Fitter")) == ["F2", "F3"]
        assert _cands(_preview(client, "?company=ROBO&retry_min=2")) == ["F2"]
        assert _cands(_preview(client, "?retry_max=0")) == ["F1"]
        recent = (datetime.utcnow() - timedelta(days=5)).strftime("%Y-%m-%d")
        assert _cands(_preview(client, f"?date_from={recent}")) == ["F1", "F2"]
        assert _cands(_preview(client, f"?date_to={recent}")) == ["F3"]
        assert _preview(client, f"?initiator_id={ini}")["total"] == 3
        chips = {c["label"]: c["value"] for c in _preview(client, f"?initiator_id={ini}&company=ROBO")["chips"]}
        assert chips["Reporting Manager"] == "Exp Ini" and chips["Company"] == "ROBO"

    def test_preview_is_capped_at_ten_rows_but_reports_the_true_total(self, client, db, world):
        aid, ini = world
        for i in range(13):
            _req(db, ini, f"Bulk {i}")
        db.session.commit()
        login(client, "expadmin@t.com")
        d = _preview(client)
        assert d["total"] == 13 and len(d["rows"]) == 10


class TestExportWorkbook:
    def test_workbook_contents_and_headers(self, client, db, world):
        aid, ini = world
        r = _req(db, ini, "Excel Person", form={"mobile_number": "9876543210", "contract_from": "2026-07-01"})
        bh = _make_user("Exp BH", "expbh@t.com", UserRole.BUSINESS_HEAD, db, companies=["RDC"])
        bh.employee_code = "BH001"
        sa = _make_user("Exp SA", "expsa@t.com", UserRole.DR_BHOON, db)
        db.session.flush()
        for actor, remark in ((bh, "ok1"), (sa, "ok2")):
            db.session.add(ApprovalAction(request_id=r.id, actor_id=actor.id, action=ApprovalActionType.APPROVED, remark=remark))
        db.session.commit()
        login(client, "expadmin@t.com")
        resp = client.get("/exports/active-employees/download")
        assert resp.status_code == 200
        assert resp.headers["Content-Type"].startswith("application/vnd.openxmlformats")
        assert 'attachment; filename="onboarding_ACTIVE_' in resp.headers["Content-Disposition"]
        assert "no-store" in resp.headers["Cache-Control"]
        ws = load_workbook(io.BytesIO(resp.data)).active
        headers = [c.value for c in ws[2]]
        for h in ("Request ID", "Status", "Initiated By", "Business / Functional Head Approval", "HR Manager Approval",
                  "Head HR Approval", "Special Approver Approval", "Submitted On", "Retry Count"):
            assert h in headers, h
        assert "BH Approval" not in headers and "HR Mgr Approval" not in headers
        row = {headers[i]: c.value for i, c in enumerate(ws[3])}
        assert row["Status"] == "Approved" and row["Initiated By"] == "Exp Ini"
        assert row["Business / Functional Head Approval"] == "Exp BH (BH001)"
        assert row["Special Approver Approval"].startswith("Exp SA")
        assert row["Associate Name (as per Aadhaar)"] == "Excel Person"
        assert row["Contract From (Date of Joining)"] == "01/07/2026"          # DD/MM/YYYY, never ISO
        assert "Generated: " in ws["A1"].value and re.search(r"Generated: \d{2}/\d{2}/\d{4} \d{2}:\d{2}", ws["A1"].value)
        assert re.fullmatch(r"\d{2}/\d{2}/\d{4}", row["Submitted On"])
        assert ws.cell(row=4, column=1).value == "Total records: 1"

    def test_download_is_audited_and_names_the_status(self, client, db, world):
        aid, ini = world
        _req(db, ini, "Aud", status=RequestStatus.PENDING_BH)
        db.session.commit()
        login(client, "expadmin@t.com")
        resp = client.get("/exports/active-employees/download?status=PENDING_BH")
        assert "onboarding_PENDING_BH_" in resp.headers["Content-Disposition"]
        row = AuditLog.query.filter_by(action_type="EXPORT_DOWNLOADED").order_by(AuditLog.id.desc()).first()
        assert '"record_count": 1' in row.detail

    @pytest.mark.parametrize("role", [UserRole.SUPER_ADMIN, UserRole.HEAD_HR, UserRole.HR_MANAGER, UserRole.DR_BHOON])
    def test_filter_page_for_each_allowed_role(self, client, db, role):
        email = f"allowed_{role.value.lower()}@t.com"
        _make_user("Allowed " + role.value, email, role, db, companies=["RDC"] if role == UserRole.HR_MANAGER else None)
        db.session.commit()
        login(client, email)
        page = client.get("/exports/active-employees").get_data(as_text=True)
        assert "Export Report" in page and "Pending Business / Functional Head" in page


class TestExportCompanyScoping:
    def test_hr_manager_only_exports_ticked_companies(self, client, db, world):
        aid, ini = world
        _req(db, ini, "Rdc Person", company="RDC")
        _req(db, ini, "Robo Person", company="ROBO")
        _make_user("Exp HRM", "exphrm@t.com", UserRole.HR_MANAGER, db, companies=["ROBO"])
        _make_user("Exp HRM0", "exphrm0@t.com", UserRole.HR_MANAGER, db)           # no ticks at all
        db.session.commit()
        login(client, "exphrm@t.com")
        d = _preview(client)
        assert d["total"] == 1 and _cands(d) == ["Robo Person"]
        ws = load_workbook(io.BytesIO(client.get("/exports/active-employees/download").data)).active
        assert ws.cell(row=4, column=1).value == "Total records: 1"
        client.get("/auth/logout")
        login(client, "exphrm0@t.com")
        assert _preview(client)["total"] == 0
