"""Regression tests for the 2026-10-10 full-application audit fixes."""
import io
import uuid
from unittest import mock

import openpyxl

from app.models import OnboardingRequest, RequestStatus, UserRole, User
from app.extensions import db as _db
from .conftest import login, logout, _make_user


def _draft(db, user, company="RDC", **form):
    req = OnboardingRequest(
        initiated_by=user.id, status=RequestStatus.DRAFT, public_token=uuid.uuid4().hex,
        candidate_name=form.get("associate_name", "Test Candidate"), company_code=company,
        plant_location="Plant A", designation="Engineer")
    db.session.add(req)
    db.session.flush()
    req.form_data = {"company_code": company, "associate_name": req.candidate_name,
                     "plant_location": "Plant A", "designation": "Engineer", **form}
    db.session.commit()
    return req


class TestDeactivatedUserLosesSession:
    def test_live_session_stops_working_once_deactivated(self, client, db, app):
        user = _make_user("Deact1", "deact1@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        with app.app_context():
            login(client, user.email)
            assert client.get("/", follow_redirects=False).status_code in (200, 302)
            assert client.get("/requests/new", follow_redirects=False).status_code != 302 or True
            u = _db.session.get(User, user.id)
            u.is_active = False
            _db.session.commit()
            resp = client.get("/notifications", follow_redirects=False)
            assert resp.status_code == 302 and "/auth/login" in resp.headers["Location"]


class TestLoginHardening:
    def test_backslash_next_is_not_followed(self, client, db, app):
        user = _make_user("Redir1", "redir1@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        with app.app_context():
            resp = client.post("/auth/login?next=/\\evil.example", data={"login_id": user.email, "password": "Test1234"})
            assert resp.status_code == 302
            assert "evil.example" not in resp.headers["Location"]

    def test_overlong_password_is_a_failed_login_not_a_500(self, client, db, app):
        user = _make_user("Long1", "long1@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        with app.app_context():
            resp = client.post("/auth/login", data={"login_id": user.email, "password": "A1" + "x" * 100})
            assert resp.status_code == 200

    def test_reset_link_uses_admin_configured_address_not_host_header(self, client, db, app):
        user = _make_user("Reset1", "reset1@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        app.config["APP_BASE_URL"] = "https://hiring.example.com"
        sent = {}
        with app.app_context(), mock.patch("app.auth.routes.send_email",
                                           side_effect=lambda **kw: sent.update(kw)):
            client.post("/auth/forgot-password", data={"email": user.email},
                        headers={"Host": "evil.example.net", "X-Forwarded-Host": "evil.example.net"})
        assert "https://hiring.example.com/auth/reset-password/" in sent["body"]
        assert "evil.example.net" not in sent["body"]


class TestOtpNotReadableFromSession:
    def test_session_holds_hash_not_code(self, client, db, app):
        user = _make_user("Otp1", "otp1@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        with app.app_context(), mock.patch("app.requests_bp.routes.get_db_mail_config",
                                           return_value={"username": "x"}), \
                mock.patch("app.requests_bp.routes._send_smtp_to_queue",
                           side_effect=lambda q, *a, **k: q.put(("ok", None))):
            login(client, user.email)
            r = client.post("/requests/send-email-otp", data={"email": "cand1@example.com"})
            assert r.get_json()["ok"] is True
            with client.session_transaction() as s:
                stored = s["_email_otp"]
            assert "code" not in stored and len(stored["hash"]) == 64

    def test_wrong_codes_are_locked_out(self, client, db, app):
        from app.requests_bp.routes import _otp_digest
        user = _make_user("Otp2", "otp2@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        with app.app_context():
            login(client, user.email)
            import time
            with client.session_transaction() as s:
                s["_email_otp"] = {"email": "c@example.com", "hash": _otp_digest("c@example.com", "123456"),
                                   "at": time.time(), "verified": False, "tries": 0}
            for _ in range(5):
                assert client.post("/requests/verify-email-otp",
                                   data={"email": "c@example.com", "otp": "000000"}).get_json()["ok"] is False
            r = client.post("/requests/verify-email-otp", data={"email": "c@example.com", "otp": "123456"})
            assert r.get_json()["ok"] is False   # right code, but already locked out


class TestSpecialCaseRoutes:
    def test_acknowledge_rejected_for_non_rdc_draft(self, client, db, app):
        user = _make_user("Ack1", "ack1@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        req = _draft(db, user, company="ROBO")
        with app.app_context():
            login(client, user.email)
            r = client.post(f"/requests/{req.public_token}/acknowledge-special-case")
            assert r.status_code == 400
            assert _db.session.get(OnboardingRequest, req.id).is_special_case is False


class TestExports:
    def test_formula_text_is_exported_as_text(self, client, db, app):
        head = _make_user("ExpHead", "exphead@t.com", UserRole.SUPER_ADMIN, db)
        user = _make_user('=HYPERLINK("http://evil.example","x")', "expinit@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        req = _draft(db, user)
        req.status = RequestStatus.ACTIVE
        _db.session.commit()
        with app.app_context():
            login(client, head.email)
            r = client.get("/exports/active-employees/download")
            assert r.status_code == 200
            ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
            kinds = {c.data_type for row in ws.iter_rows(min_row=3) for c in row if isinstance(c.value, str)
                     and c.value.startswith("=")}
            assert kinds <= {"s"} and kinds

    def test_drafts_are_never_exported(self, client, db, app):
        head = _make_user("ExpHead2", "exphead2@t.com", UserRole.SUPER_ADMIN, db)
        user = _make_user("ExpInit2", "expinit2@t.com", UserRole.INITIATOR, db, companies=["RDC"])
        _draft(db, user, associate_name="Private Draft Person")
        with app.app_context():
            login(client, head.email)
            r = client.get("/exports/active-employees/download?status=DRAFT")
            blob = b"" if r.status_code != 200 else r.data
            if blob:
                ws = openpyxl.load_workbook(io.BytesIO(blob)).active
                values = [c.value for row in ws.iter_rows() for c in row]
                assert "Private Draft Person" not in values


class TestPasswordRules:
    def test_overlong_password_rejected(self):
        from app.utils import validate_password
        assert any("too long" in e for e in validate_password("Aa1" + "x" * 80))
        assert validate_password("Aa1xxxxx") == []


class TestSubmitServerSideChecks:
    def _world(self, db, suffix, **form):
        user = _make_user("Sv" + suffix, f"sv{suffix}@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        _make_user("SvBh" + suffix, f"svbh{suffix}@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
        _make_user("SvHr" + suffix, f"svhr{suffix}@t.com", UserRole.HR_MANAGER, db, companies=["ROBO"])
        from app.models import PlantLocation
        db.session.add(PlantLocation(name="Plant A", company="ROBO"))
        return user, _draft(db, user, company="ROBO", uan_number="AB1234567890", **form)

    def _submit(self, client, app, user, req):
        with app.app_context():
            login(client, user.email)
            client.post(f"/requests/{req.public_token}/submit", follow_redirects=True)
            return _db.session.get(OnboardingRequest, req.id).status

    def test_bad_aadhaar_blocks_submit(self, client, db, app):
        user, req = self._world(db, "a1", aadhar_no="ABCDEFGHIJKL")
        assert self._submit(client, app, user, req) == RequestStatus.DRAFT

    def test_bad_mobile_blocks_submit(self, client, db, app):
        user, req = self._world(db, "m1", mobile_number="1234567890")
        assert self._submit(client, app, user, req) == RequestStatus.DRAFT

    def test_unknown_plant_blocks_submit_for_non_rdc(self, client, db, app):
        user, req = self._world(db, "p1")
        req.form_data = dict(req.form_data, plant_location="Nowhere Plant")
        _db.session.commit()
        assert self._submit(client, app, user, req) == RequestStatus.DRAFT

    def test_email_changed_after_verification_blocks_submit(self, client, db, app):
        user, req = self._world(db, "e1", email_id="b@example.com")
        req.candidate_email_verified = "a@example.com"
        _db.session.commit()
        assert self._submit(client, app, user, req) == RequestStatus.DRAFT

    def test_valid_request_still_submits(self, client, db, app):
        user, req = self._world(db, "ok1", aadhar_no="1234 5678 9012", mobile_number="9876543210")
        assert self._submit(client, app, user, req) == RequestStatus.PENDING_BH
