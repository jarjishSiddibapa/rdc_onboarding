"""
Request event ledger (2026-10-14): every step in a request's life is stored in request_events with the
exact time, the actor, and how long the request had been waiting - the data for later response-time
analysis - and the request page's timeline shows the same times and durations.
"""
from app.timeutil import now_ist
import uuid
from datetime import datetime, timedelta

from app.extensions import db as _db
from app.models import (
    ApprovalAction, OnboardingRequest, PlantLocation, RequestEvent, RequestStatus, UserRole,
)
from app.utils import format_duration, record_request_event
from .conftest import login, logout, _make_user

T0 = datetime(2026, 10, 14, 4, 0, 0)


def _req(db, initiator, status=RequestStatus.DRAFT, company="ROBO", created_at=T0):
    r = OnboardingRequest(initiated_by=initiator.id, status=status, public_token=uuid.uuid4().hex,
                          candidate_name="Event Candidate", company_code=company, plant_location="Plant A",
                          designation="Engineer", created_at=created_at)
    db.session.add(r)
    db.session.flush()
    r.form_data = {"company_code": company, "associate_name": "Event Candidate", "plant_location": "Plant A",
                   "designation": "Engineer", "uan_number": "AB1234567890"}
    if company != "RDC" and not PlantLocation.query.filter_by(name="Plant A", company=company).first():
        db.session.add(PlantLocation(name="Plant A", company=company))
    db.session.flush()
    return r


class TestFormatDuration:
    def test_formats(self):
        assert format_duration(None) == "—"
        assert format_duration(0) == "0 sec"
        assert format_duration(45) == "45 sec"
        assert format_duration(125) == "2 min 5 sec"
        assert format_duration(600) == "10 min"
        assert format_duration(3 * 3600 + 20 * 60) == "3 h 20 min"
        assert format_duration(2 * 86400 + 4 * 3600) == "2 d 4 h"
        assert format_duration(timedelta(hours=5)) == "5 h"
        assert format_duration(-30) == "0 sec"


class TestRecorder:
    def test_durations_and_cycle_follow_the_hand_offs(self, client, db, app):
        with app.app_context():
            init = _make_user("EvInit1", "evinit1@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
            bh = _make_user("EvBh1", "evbh1@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
            req = _req(db, init)
            ev0 = record_request_event(req, "DRAFT_CREATED", actor=init, at=T0)
            sub = record_request_event(req, "SUBMITTED", actor=init, at=T0 + timedelta(seconds=90))
            ack = record_request_event(req, "SPECIAL_CASE_ACKNOWLEDGED", actor=init, at=T0 + timedelta(seconds=100))
            app1 = record_request_event(req, "APPROVED", actor=bh, from_status="PENDING_BH",
                                        to_status="PENDING_HR_MANAGER", remark="ok",
                                        at=T0 + timedelta(seconds=90 + 3600))
            rej = record_request_event(req, "REJECTED", actor=bh, at=T0 + timedelta(seconds=90 + 3600 + 600))
            db.session.flush()
            assert ev0.duration_seconds is None and ev0.waiting_since is None
            assert sub.waiting_since == T0 and sub.duration_seconds == 90          # time spent drafting
            assert app1.waiting_since == sub.occurred_at and app1.duration_seconds == 3600
            assert app1.cycle_seconds == 3600                                       # since this attempt was submitted
            # the acknowledgement is not a hand-off: the next decision still measures from the hand-off
            assert rej.waiting_since == app1.occurred_at and rej.duration_seconds == 600
            assert rej.cycle_seconds == 3600 + 600
            assert (app1.actor_name, app1.actor_role, app1.remark) == ("EvBh1", "BUSINESS_HEAD", "ok")
            assert app1.company_code == "ROBO" and app1.attempt_no == 1

    def test_resubmission_measures_the_time_to_correct_and_counts_the_attempt(self, client, db, app):
        with app.app_context():
            init = _make_user("EvInit2", "evinit2@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
            req = _req(db, init)
            record_request_event(req, "SUBMITTED", actor=init, at=T0 + timedelta(minutes=1))
            rej = record_request_event(req, "REJECTED", actor=init, at=T0 + timedelta(minutes=61))
            req.retry_count = 1
            re_ = record_request_event(req, "RESUBMITTED", actor=init, at=T0 + timedelta(minutes=181))
            appr = record_request_event(req, "APPROVED", actor=init, at=T0 + timedelta(minutes=191))
            db.session.flush()
            assert re_.duration_seconds == 120 * 60 and re_.attempt_no == 2
            assert appr.cycle_seconds == 600          # measured from the resubmission, not the first submit

    def test_system_event_has_no_actor_and_never_raises(self, client, db, app):
        with app.app_context():
            init = _make_user("EvInit3", "evinit3@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
            req = _req(db, init)
            record_request_event(req, "APPROVED", actor=init, at=T0 + timedelta(minutes=5))
            push = record_request_event(req, "TRUEIN_PUSHED", via="system", at=T0 + timedelta(minutes=5, seconds=7))
            db.session.flush()
            assert push.actor_id is None and push.actor_name is None and push.via == "system"
            assert push.duration_seconds == 7         # seconds after the approval that triggered it


class TestRoutesWriteEvents:
    def _chain(self, db):
        init = _make_user("RtInit", "rtinit@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        bh = _make_user("RtBh", "rtbh@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
        hrm = _make_user("RtHrm", "rthrm@t.com", UserRole.HR_MANAGER, db, companies=["ROBO"])
        req = _req(db, init, created_at=now_ist() - timedelta(minutes=30))
        db.session.commit()
        return init.email, bh.email, hrm.email, req.public_token, req.id

    def test_submit_approve_reject_resubmit_are_all_recorded(self, client, db, app):
        init_e, bh_e, hrm_e, token, rid = self._chain(db)
        with app.app_context():
            login(client, init_e)
            client.post(f"/requests/{token}/submit", follow_redirects=True)
            logout(client)
            login(client, bh_e)
            client.post(f"/requests/{token}/approve", data={"remark": "bh ok"}, follow_redirects=True)
            logout(client)
            login(client, hrm_e)
            client.post(f"/requests/{token}/reject", data={"remark": "fix the address"}, follow_redirects=True)
            logout(client)
            login(client, init_e)
            client.post(f"/requests/{token}/resubmit", follow_redirects=True)

            evs = RequestEvent.query.filter_by(request_id=rid).order_by(RequestEvent.id).all()
            assert [e.event_type for e in evs] == ["SUBMITTED", "APPROVED", "REJECTED", "RESUBMITTED"]
            sub, approved, rejected, resub = evs
            # the drafting time is real (the draft was created 30 minutes before it was submitted)
            assert 29 * 60 <= sub.duration_seconds <= 31 * 60
            assert (approved.from_status, approved.to_status) == ("PENDING_BH", "PENDING_HR_MANAGER")
            assert (rejected.from_status, rejected.to_status) == ("PENDING_HR_MANAGER", "REJECTED_HRM")
            assert approved.actor_role == "BUSINESS_HEAD" and rejected.actor_role == "HR_MANAGER"
            assert rejected.waiting_since == approved.occurred_at
            assert resub.attempt_no == 2 and resub.waiting_since == rejected.occurred_at
            assert all(e.via == "web" for e in evs)
            # the event and the approval row are the same instant and linked
            for e in (approved, rejected):
                act = _db.session.get(ApprovalAction, e.approval_action_id)
                assert act.acted_at == e.occurred_at and act.actor_id == e.actor_id and act.remark == e.remark

    def test_timeline_shows_exact_times_and_how_long_each_step_took(self, client, db, app):
        init_e, bh_e, hrm_e, token, rid = self._chain(db)
        with app.app_context():
            login(client, init_e)
            client.post(f"/requests/{token}/submit", follow_redirects=True)
            logout(client)
            login(client, bh_e)
            client.post(f"/requests/{token}/approve", data={"remark": "bh ok"}, follow_redirects=True)
            logout(client)
            login(client, init_e)
            html = client.get(f"/requests/{token}").get_data(as_text=True)
        assert "Responded in" in html            # how long the Business Head took
        assert "Draft took" in html              # how long the draft took before it was submitted
        assert "Waiting since" in html and "so far" in html     # the HR Manager stage is pending right now
        assert "Activity log (2)" in html        # submitted, approved (this draft was built directly, not via /new)
        assert "Submitted for approval" in html
        # exact time, including seconds, e.g. 14/10/2026, 09:30:05
        import re
        assert re.search(r"\d\d/\d\d/\d{4}, \d\d:\d\d:\d\d", html)

    def test_deleting_a_draft_is_recorded(self, client, db, app):
        init = _make_user("RtInit2", "rtinit2@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        req = _req(db, init)
        db.session.commit()
        token, rid, email = req.public_token, req.id, init.email
        with app.app_context():
            login(client, email)
            client.post(f"/requests/{token}/delete", follow_redirects=True)
            types = [e.event_type for e in RequestEvent.query.filter_by(request_id=rid)]
        assert types == ["DRAFT_DELETED"]


class TestDraftAndPushEvents:
    def test_opening_a_new_request_records_draft_created_once(self, client, db, app):
        init = _make_user("EvNew", "evnew@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        db.session.commit()
        email, uid = init.email, init.id
        with app.app_context():
            login(client, email)
            client.get("/requests/new")
            client.get("/requests/new")       # reuses the blank draft - no second DRAFT_CREATED
            evs = RequestEvent.query.filter_by(event_type="DRAFT_CREATED").all()
            assert len(evs) == 1 and evs[0].actor_id == uid and evs[0].to_status == "DRAFT"
            assert evs[0].duration_seconds is None

    def test_every_truein_attempt_is_recorded(self, client, db, app):
        from app.integrations.truein import _write_push_log
        with app.app_context():
            init = _make_user("EvPush", "evpush@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
            req = _req(db, init, status=RequestStatus.ACTIVE)
            record_request_event(req, "APPROVED", actor=init, at=now_ist() - timedelta(seconds=30))
            req.truein_retry_count = 1
            _write_push_log(_db, req, {"success": False, "message": "Govt ID already exist", "empId": "X1"}, triggered_by="auto")
            req.truein_retry_count = 2
            _write_push_log(_db, req, {"success": True, "message": "ok", "empId": "X1"}, triggered_by="auto")
            _db.session.flush()
            evs = RequestEvent.query.filter(RequestEvent.request_id == req.id,
                                            RequestEvent.event_type.like("TRUEIN%")).order_by(RequestEvent.id).all()
            assert [e.event_type for e in evs] == ["TRUEIN_PUSH_FAILED", "TRUEIN_PUSHED"]
            assert evs[0].remark == "Govt ID already exist" and evs[0].via == "system" and evs[0].actor_id is None
            assert evs[1].duration_seconds is not None and evs[1].duration_seconds >= 0


class TestActionTimesExport:
    def _setup(self, db, app, client):
        init = _make_user("ExInit", "exinit@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        bh = _make_user("ExBh", "exbh@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
        hhr = _make_user("ExHhr", "exhhr@t.com", UserRole.HEAD_HR, db)
        req = _req(db, init, status=RequestStatus.PENDING_BH)
        db.session.commit()
        for kind, who, mins in (("SUBMITTED", init, 2), ("APPROVED", bh, 62), ("APPROVED", bh, 122)):
            record_request_event(req, kind, actor=who, from_status="PENDING_BH", to_status="PENDING_HR_MANAGER",
                                 at=T0 + timedelta(minutes=mins))
        db.session.commit()
        return hhr.email, bh.email

    def test_head_hr_downloads_log_and_per_role_averages(self, client, db, app):
        hhr_email, _ = self._setup(db, app, client)
        with app.app_context():
            login(client, hhr_email)
            resp = client.get("/exports/action-times/download")
        assert resp.status_code == 200
        assert resp.headers["Content-Type"].startswith("application/vnd.openxmlformats")
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.data))
        assert wb.sheetnames == ["Action log", "Response time by role"]
        log = [r for r in wb["Action log"].iter_rows(min_row=3, values_only=True)]
        assert [r[6] for r in log] == ["Submitted for approval", "Approved", "Approved"]
        assert log[1][11] == "14/10/2026 05:02:00"          # stored and shown as Indian time
        stats = [r for r in wb["Response time by role"].iter_rows(min_row=3, values_only=True)]
        all_co = [r for r in stats if r[1] == "All companies"][0]
        assert all_co[0] == "Business / Functional Head" and all_co[2] == 2
        assert all_co[6] == 60.0     # each decision came an hour after the previous hand-off

    def test_initiator_cannot_download(self, client, db, app):
        init = _make_user("ExInit2", "exinit2@t.com", UserRole.INITIATOR, db, companies=["ROBO"])
        db.session.commit()
        with app.app_context():
            login(client, init.email)
            assert client.get("/exports/action-times/download").status_code == 403
