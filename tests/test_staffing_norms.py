"""
Unit tests for app/services/staffing_norms.py — the RDC hiring-gate logic.

Everything here is DB-only (StaffingSnapshot rows seeded directly) plus a
mocked dvt.get_average_plant_volume()/get_average_cluster_total_volume() — no live API
calls, matching how check_rdc_staffing_gate() actually behaves at request
time (it never calls ZingHR/Truein/DVT synchronously except the one DVT
volume lookup, which we mock here).
"""
from datetime import datetime
from unittest.mock import patch

import pytest

from app.models import (
    OnboardingRequest, RequestStatus, Designation, NormRoleCategory, NormTier,
    NormRequirement, NormScope, NormSheet, NormRequirementType,
    PlantDvtMapping, ClusterNameMapping, StaffingSnapshot, MatchConfidence,
    BusinessHeadRegion, UserRole, EmployeeLocationSnapshot, ExternalDesignationSource,
    PlantLocation,
)
from app.services import staffing_norms
from .conftest import _make_user, login


def _make_request(db, initiator, designation, plant_location):
    req = OnboardingRequest(initiated_by=initiator.id, status=RequestStatus.DRAFT)
    db.session.add(req)
    db.session.flush()
    req.form_data = {
        "company_code": "RDC",
        "designation": designation,
        "plant_location": plant_location,
    }
    db.session.flush()
    return req


def _seed_plant_norm(db, role_name="Batchers/Production Officer", tier_key="3000_5000",
                      min_v=3000, max_v=5000, req_type=NormRequirementType.FIXED,
                      fixed_count=2, rate_per_unit=None, unit_volume=None):
    cat = NormRoleCategory(name=role_name, scope=NormScope.PLANT, sheet=NormSheet.SHEET1)
    db.session.add(cat)
    db.session.flush()
    tier = NormTier(sheet=NormSheet.SHEET1, scope=NormScope.PLANT, tier_key=tier_key,
                     tier_label=tier_key, min_value=min_v, max_value=max_v)
    db.session.add(tier)
    db.session.flush()
    req = NormRequirement(tier_id=tier.id, norm_role_category_id=cat.id,
                           requirement_type=req_type, fixed_count=fixed_count,
                           rate_per_unit=rate_per_unit, unit_volume=unit_volume)
    db.session.add(req)
    db.session.flush()
    return cat, tier, req


class TestNotCoveredAndUnmapped:
    def test_designation_not_in_master_allows(self, db, initiator):
        req = _make_request(db, initiator, "Some Unknown Role", "PlantX")
        result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "not_covered"

    def test_designation_with_no_norm_category_allows(self, db, initiator):
        desig = Designation(name="HR Executive", norm_category_id=None)
        db.session.add(desig)
        db.session.flush()
        req = _make_request(db, initiator, "HR Executive", "PlantX")
        result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "not_covered"

    def test_plant_not_mapped_allows(self, db, initiator):
        cat, tier, norm_req = _seed_plant_norm(db)
        desig = Designation(name="Batcher", norm_category_id=cat.id)
        db.session.add(desig)
        db.session.flush()
        req = _make_request(db, initiator, "Batcher", "Unmapped Plant")
        result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "plant_not_mapped"


class TestPlantScopeFixed:
    def _setup(self, db, initiator, current_headcount):
        cat, tier, norm_req = _seed_plant_norm(db, fixed_count=2)
        desig = Designation(name="Batcher", norm_category_id=cat.id)
        db.session.add(desig)
        plant_map = PlantDvtMapping(plant_location_name="PlantX", dvt_plant_code="PX1")
        db.session.add(plant_map)
        db.session.flush()
        db.session.add(StaffingSnapshot(
            scope=NormScope.PLANT, location_key="PlantX", norm_role_category_id=cat.id,
            current_headcount=current_headcount, zinghr_count=current_headcount, truein_count=0,
        ))
        db.session.flush()
        return _make_request(db, initiator, "Batcher", "PlantX")

    def test_under_norm_allows(self, db, initiator):
        req = self._setup(db, initiator, current_headcount=1)
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", return_value=4000.0):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "ok"
        assert result["details"]["current_headcount"] == 1
        assert result["details"]["allowed_headcount"] == 2
        assert result["details"]["would_be_headcount"] == 2

    def test_at_norm_blocks(self, db, initiator):
        req = self._setup(db, initiator, current_headcount=2)
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", return_value=4000.0):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is False
        assert result["reason"] == "at_or_over_norm"
        assert result["details"]["would_be_headcount"] == 3
        assert result["details"]["allowed_headcount"] == 2

    def test_no_snapshot_yet_allows(self, db, initiator):
        cat, tier, norm_req = _seed_plant_norm(db, fixed_count=2)
        desig = Designation(name="Batcher", norm_category_id=cat.id)
        db.session.add(desig)
        plant_map = PlantDvtMapping(plant_location_name="PlantX", dvt_plant_code="PX1")
        db.session.add(plant_map)
        db.session.flush()
        req = _make_request(db, initiator, "Batcher", "PlantX")
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", return_value=4000.0):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "no_snapshot_yet"


class TestAdjacentTierBoundaries:
    """
    Regression coverage for the 2026-09-21 "< 1500 m3" tier addition:
    _find_tier() (both the copy in staffing_norms.py and the one in
    headcount.py) scans NormTier rows with no ORDER BY and returns the
    FIRST range match. Two tiers must never have overlapping
    [min_value, max_value) ranges for the same scope/sheet, or which one
    "wins" depends on unspecified row order. These tests lock in that the
    lower tier's upper bound is exclusive and the next tier's lower bound
    is inclusive, using non-overlapping ranges exactly like the real
    LT_1500 (None, 1500) / LT_3000 (1500, 3000) pair.
    """

    def _seed_two_adjacent_tiers(self, db):
        cat = NormRoleCategory(name="Batcher Role", scope=NormScope.PLANT, sheet=NormSheet.SHEET1)
        db.session.add(cat)
        db.session.flush()
        lo_tier = NormTier(sheet=NormSheet.SHEET1, scope=NormScope.PLANT, tier_key="LT_1500",
                            tier_label="< 1500 m3", min_value=None, max_value=1500)
        hi_tier = NormTier(sheet=NormSheet.SHEET1, scope=NormScope.PLANT, tier_key="1500_3000",
                            tier_label="1500-3000 m3", min_value=1500, max_value=3000)
        db.session.add(lo_tier)
        db.session.add(hi_tier)
        db.session.flush()
        db.session.add(NormRequirement(tier_id=lo_tier.id, norm_role_category_id=cat.id,
                                        requirement_type=NormRequirementType.FIXED, fixed_count=1))
        db.session.add(NormRequirement(tier_id=hi_tier.id, norm_role_category_id=cat.id,
                                        requirement_type=NormRequirementType.FIXED, fixed_count=2))
        db.session.flush()
        return cat, lo_tier, hi_tier

    def test_find_tier_boundary_is_exclusive_below_inclusive_above(self, db):
        cat, lo_tier, hi_tier = self._seed_two_adjacent_tiers(db)
        assert staffing_norms._find_tier(NormScope.PLANT, NormSheet.SHEET1, 1499).id == lo_tier.id
        assert staffing_norms._find_tier(NormScope.PLANT, NormSheet.SHEET1, 1500).id == hi_tier.id
        assert staffing_norms._find_tier(NormScope.PLANT, NormSheet.SHEET1, 2999).id == hi_tier.id

    def test_gate_uses_correct_tier_allowed_headcount_at_boundary(self, db, initiator):
        cat, lo_tier, hi_tier = self._seed_two_adjacent_tiers(db)
        desig = Designation(name="Batcher", norm_category_id=cat.id)
        db.session.add(desig)
        plant_map = PlantDvtMapping(plant_location_name="PlantX", dvt_plant_code="PX1")
        db.session.add(plant_map)
        db.session.flush()
        db.session.add(StaffingSnapshot(
            scope=NormScope.PLANT, location_key="PlantX", norm_role_category_id=cat.id,
            current_headcount=0, zinghr_count=0, truein_count=0,
        ))
        db.session.flush()
        req = _make_request(db, initiator, "Batcher", "PlantX")
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", return_value=1200.0):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["details"]["allowed_headcount"] == 1
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", return_value=1500.0):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["details"]["allowed_headcount"] == 2


class TestRateBasedRounding:
    def test_rounds_to_nearest(self, db, initiator):
        # tier is 3000-5000 m3 (see _seed_plant_norm defaults); 1 per 900 m3,
        # volume=4200 -> 4200/900 = 4.666... -> rounds to 5 (nearest, not floor)
        cat, tier, norm_req = _seed_plant_norm(
            db, role_name="FTs/LT/TO", req_type=NormRequirementType.RATE_PER_VOLUME,
            fixed_count=None, rate_per_unit=1, unit_volume=900,
        )
        desig = Designation(name="TM Driver", norm_category_id=cat.id)
        db.session.add(desig)
        plant_map = PlantDvtMapping(plant_location_name="PlantX", dvt_plant_code="PX1")
        db.session.add(plant_map)
        db.session.flush()
        db.session.add(StaffingSnapshot(
            scope=NormScope.PLANT, location_key="PlantX", norm_role_category_id=cat.id,
            current_headcount=4, zinghr_count=4, truein_count=0,
        ))
        db.session.flush()
        req = _make_request(db, initiator, "TM Driver", "PlantX")
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", return_value=4200.0):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["details"]["allowed_headcount"] == 5  # rounds up from 4.6667, not floors to 4
        assert result["allowed"] is True  # 4 current + 1 = 5 <= 5 allowed

    def test_compute_allowed_headcount_rounds_down_case(self):
        # 1 per 900, volume=2200 -> 2.44 -> rounds to nearest = 2
        req = NormRequirement(requirement_type=NormRequirementType.RATE_PER_VOLUME,
                               rate_per_unit=1, unit_volume=900)
        assert staffing_norms.compute_allowed_headcount(req, 2200.0) == 2


class TestClusterScopePerBusinessHead:
    def test_per_business_head_skips_gate(self, db, initiator):
        cluster = ClusterNameMapping(canonical_cluster_name="DELHI NCR")
        db.session.add(cluster)
        db.session.flush()
        cat = NormRoleCategory(name="EA", scope=NormScope.CLUSTER, sheet=NormSheet.SHEET1)
        db.session.add(cat)
        db.session.flush()
        tier = NormTier(sheet=NormSheet.SHEET1, scope=NormScope.CLUSTER, tier_key="GT_4_PLANTS",
                         tier_label="> 4 Plants", min_value=5, max_value=None)
        db.session.add(tier)
        db.session.flush()
        norm_req = NormRequirement(tier_id=tier.id, norm_role_category_id=cat.id,
                                    requirement_type=NormRequirementType.PER_BUSINESS_HEAD)
        db.session.add(norm_req)
        desig = Designation(name="EA to BH", norm_category_id=cat.id)
        db.session.add(desig)
        # 5 plants in the cluster so it resolves to the >4 tier
        for i in range(5):
            db.session.add(PlantDvtMapping(plant_location_name=f"Plant{i}", dvt_plant_code=f"P{i}", cluster_id=cluster.id))
        db.session.flush()
        req = _make_request(db, initiator, "EA to BH", "Plant0")
        result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "per_business_head_not_supported"


class TestErrorHandling:
    def test_dvt_exception_with_no_cached_volume_fails_open(self, db, initiator):
        """
        DVT raises AND there's no prior StaffingSnapshot volume to fall back
        to (e.g. right after a fresh deploy, before any snapshot has ever
        run) — genuinely no volume data available, so this fails open with
        the specific "no_volume_data" reason, same as a clean DVT None
        response would. See test_dvt_exception_falls_back_to_last_known_volume
        for the case where a fallback IS available.
        """
        cat, tier, norm_req = _seed_plant_norm(db)
        desig = Designation(name="Batcher", norm_category_id=cat.id)
        db.session.add(desig)
        plant_map = PlantDvtMapping(plant_location_name="PlantX", dvt_plant_code="PX1")
        db.session.add(plant_map)
        db.session.flush()
        req = _make_request(db, initiator, "Batcher", "PlantX")
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", side_effect=RuntimeError("DVT down")):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is True
        assert result["reason"] == "no_volume_data"

    def test_dvt_exception_falls_back_to_last_known_volume(self, db, initiator):
        """
        A transient DVT outage (confirmed happening for real against the
        live DVT server) must not silently skip a genuine capacity block —
        check_rdc_staffing_gate() should fall back to the last known-good
        volume from a prior StaffingSnapshot run rather than failing open.
        """
        cat, tier, norm_req = _seed_plant_norm(db, fixed_count=1)
        desig = Designation(name="Batcher", norm_category_id=cat.id)
        db.session.add(desig)
        plant_map = PlantDvtMapping(plant_location_name="PlantX", dvt_plant_code="PX1")
        db.session.add(plant_map)
        db.session.add(StaffingSnapshot(
            scope=NormScope.PLANT, location_key="PlantX", norm_role_category_id=cat.id,
            current_headcount=1, zinghr_count=1, truein_count=0, deduped_count=0,
            unclassified_count=0, allowed_headcount=1, tier_label=tier.tier_label,
            volume_used=4000.0, production_volume=4000.0,
        ))
        db.session.flush()
        req = _make_request(db, initiator, "Batcher", "PlantX")
        with patch("app.services.staffing_norms.dvt.get_average_plant_volume", side_effect=RuntimeError("DVT down")):
            result = staffing_norms.check_rdc_staffing_gate(req.form_data)
        assert result["allowed"] is False
        assert result["reason"] == "at_or_over_norm"
        assert result["details"]["volume_used"] == 4000.0


class TestStaffingStatusDownload:
    """
    Excel download of the RDC Staffing Status dashboard (added 2026-09-10) —
    same data source and Business-Head region scoping as the on-screen
    staffing_status()/staffing_status_cluster()/staffing_status_plant() views.
    """

    def _seed(self, db):
        cluster = ClusterNameMapping(canonical_cluster_name="Bangalore")
        db.session.add(cluster)
        db.session.flush()
        cat, tier, norm_req = _seed_plant_norm(db, fixed_count=2)
        plant_map = PlantDvtMapping(
            plant_location_name="PlantX", dvt_plant_code="PX1",
            cluster_id=cluster.id, match_confidence=MatchConfidence.AUTO_EXACT,
        )
        db.session.add(plant_map)
        db.session.flush()
        # get_snapshot_rows_for_location() matches on computed_at == MAX(computed_at)
        # across ALL snapshot rows — both rows of a real "run" share one timestamp,
        # so seed them with the same explicit value rather than relying on two
        # separate datetime.utcnow() defaults (which can differ by microseconds
        # and silently exclude one of the two from "the latest run").
        run_time = datetime.utcnow()
        db.session.add(StaffingSnapshot(
            scope=NormScope.PLANT, location_key="PlantX", norm_role_category_id=cat.id,
            current_headcount=1, allowed_headcount=2, tier_label=tier.tier_label,
            production_volume=4000.0, can_hire=True, computed_at=run_time,
        ))
        db.session.add(StaffingSnapshot(
            scope=NormScope.CLUSTER, location_key="Bangalore", norm_role_category_id=cat.id,
            current_headcount=1, allowed_headcount=2, tier_label=tier.tier_label,
            production_volume=4000.0, can_hire=True, computed_at=run_time,
        ))
        # One plant-level employee and one cluster-only (no specific plant)
        # employee, sharing the same run_time — see the shared-timestamp note
        # in get_snapshot_rows_for_location()/EmployeeLocationSnapshot reads.
        db.session.add(EmployeeLocationSnapshot(
            source=ExternalDesignationSource.ZINGHR, employee_code="E001",
            employee_name="Plant Employee", designation="Operator", department="Technical",
            date_of_joining="12 May 2020", plant_location_key="PlantX",
            cluster_location_key="Bangalore", computed_at=run_time,
        ))
        db.session.add(EmployeeLocationSnapshot(
            source=ExternalDesignationSource.TRUEIN, employee_code="E002",
            employee_name="Cluster Only Employee", designation="Accountant", department="Accounts",
            date_of_joining="01 Jan 2021", plant_location_key=None,
            cluster_location_key="Bangalore", computed_at=run_time,
        ))
        db.session.flush()
        return cluster

    def test_super_admin_download_has_both_sheets_with_expected_rows(self, client, db, app):
        from io import BytesIO
        from openpyxl import load_workbook
        self._seed(db)
        admin = _make_user("DlAdmin", "dladmin@t.com", UserRole.SUPER_ADMIN, db)
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status/download")
        assert resp.status_code == 200
        assert resp.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        wb = load_workbook(BytesIO(resp.data))
        assert wb.sheetnames == ["By Cluster", "By Plant", "By Employees"]
        cluster_rows = list(wb["By Cluster"].iter_rows(min_row=2, values_only=True))
        plant_rows = list(wb["By Plant"].iter_rows(min_row=2, values_only=True))
        employee_rows = list(wb["By Employees"].iter_rows(min_row=2, values_only=True))
        assert len(cluster_rows) == 1
        assert cluster_rows[0][0] == "Bangalore"
        assert len(plant_rows) == 1
        assert plant_rows[0][0] == "Bangalore"
        assert wb["By Employees"]["A1":"H1"][0][0].value == "Cluster"
        assert wb["By Employees"]["A1":"H1"][0][1].value == "Plant"
        assert len(employee_rows) == 2
        by_code = {r[3]: r for r in employee_rows}
        plant_emp = by_code["E001"]
        assert plant_emp[0] == "Bangalore" and plant_emp[1] == "PlantX"
        assert plant_emp[2] == "Plant Employee" and plant_emp[4] == "Operator"
        cluster_only_emp = by_code["E002"]
        # openpyxl reads a cell written with value="" back as None, not "" —
        # both mean "blank Plant column" for a cluster-only employee.
        assert cluster_only_emp[0] == "Bangalore" and not cluster_only_emp[1]
        assert cluster_only_emp[2] == "Cluster Only Employee"

    def test_business_head_only_sees_own_region(self, client, db, app):
        self._seed(db)
        other_cluster = ClusterNameMapping(canonical_cluster_name="Chennai")
        db.session.add(other_cluster)
        db.session.flush()
        bh = _make_user("DlBh", "dlbh@t.com", UserRole.BUSINESS_HEAD, db)
        db.session.add(BusinessHeadRegion(business_head_id=bh.id, cluster_id=other_cluster.id))
        db.session.commit()
        with app.app_context():
            login(client, bh.email)
            resp = client.get("/requests/staffing-status/download")
        assert resp.status_code == 200
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(resp.data))
        cluster_rows = list(wb["By Cluster"].iter_rows(min_row=2, values_only=True))
        # BH is scoped to Chennai only, which has no snapshot rows -> no rows at all
        # (definitely not Bangalore's).
        assert all(r[0] != "Bangalore" for r in cluster_rows)


class TestStaffingStatusRegionEmployeePanel:
    """
    Coverage for the 2026-09-10 addition: the region (cluster) staffing page
    now shows every employee resolved anywhere in that region — every plant
    plus cluster-only staff — not just the narrower "cluster-only, not tied
    to a specific plant" table that was already there.
    """

    def _seed(self, db):
        cluster = ClusterNameMapping(canonical_cluster_name="Bangalore")
        db.session.add(cluster)
        db.session.flush()
        plant_map = PlantDvtMapping(
            plant_location_name="PlantX", dvt_plant_code="PX1",
            cluster_id=cluster.id, match_confidence=MatchConfidence.AUTO_EXACT,
        )
        db.session.add(plant_map)
        db.session.flush()
        db.session.add(EmployeeLocationSnapshot(
            source=ExternalDesignationSource.TRUEIN, employee_code="T001",
            employee_name="Region Panel Test Employee", designation="Operator",
            plant_location_key="PlantX", cluster_location_key="Bangalore",
            computed_at=datetime.utcnow(),
        ))
        db.session.flush()
        return cluster

    def test_plant_level_employee_appears_in_region_panel(self, client, db, app):
        self._seed(db)
        admin = _make_user("RegPanelAdmin", "regpaneladmin@t.com", UserRole.SUPER_ADMIN, db)
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status/cluster/Bangalore")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert "Region Panel Test Employee" in html
        assert "Every Employee in Bangalore" in html

    def test_full_directory_link_pre_filters_by_region(self, client, db, app):
        self._seed(db)
        admin = _make_user("RegPanelAdmin2", "regpaneladmin2@t.com", UserRole.SUPER_ADMIN, db)
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status/employees?cluster=Bangalore")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert "Region Panel Test Employee" in html


class TestStaffingStatusCompanyTabs:
    """
    Coverage for the 2026-09-15 multi-company fix: the Staffing Status page
    is now a 3-tab page (RDC / Ultrafine / ROBO). RDC's own tab content is
    completely unchanged (still tested by every other class in this file);
    these confirm the new tabs render and the Ultrafine/ROBO plant-detail
    route works and is properly isolated from RDC.
    """

    def test_page_renders_all_three_company_tabs(self, client, db, app):
        admin = _make_user("CompanyTabAdmin", "companytabadmin@t.com", UserRole.SUPER_ADMIN, db)
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert 'id="ctab-btn-rdc"' in html
        assert 'id="ctab-btn-ultrafine"' in html
        assert 'id="ctab-btn-robo"' in html

    def test_ultrafine_tab_shows_its_plants_with_headcount(self, client, db, app):
        admin = _make_user("CompanyTabAdmin2", "companytabadmin2@t.com", UserRole.SUPER_ADMIN, db)
        db.session.add(PlantLocation(name="UF Tab Plant", company="Ultrafine", is_active=True))
        db.session.commit()
        db.session.add(EmployeeLocationSnapshot(
            computed_at=datetime.utcnow(), source=ExternalDesignationSource.ZINGHR,
            employee_code="UFTAB1", employee_name="UF Tab Employee",
            plant_location_key="UF Tab Plant", company="Ultrafine",
        ))
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status")
        html = resp.get_data(as_text=True)
        assert "UF Tab Plant" in html

    def test_company_plant_detail_shows_only_that_companys_employees(self, client, db, app):
        admin = _make_user("CompanyTabAdmin3", "companytabadmin3@t.com", UserRole.SUPER_ADMIN, db)
        db.session.add_all([
            PlantLocation(name="ROBO Detail Plant", company="ROBO", is_active=True),
            PlantLocation(name="UF Other Plant", company="Ultrafine", is_active=True),
        ])
        db.session.commit()
        shared_now = datetime.utcnow()
        db.session.add_all([
            EmployeeLocationSnapshot(
                computed_at=shared_now, source=ExternalDesignationSource.ZINGHR,
                employee_code="ROBODET1", employee_name="Robo Detail Employee",
                plant_location_key="ROBO Detail Plant", company="ROBO",
            ),
            EmployeeLocationSnapshot(
                computed_at=shared_now, source=ExternalDesignationSource.ZINGHR,
                employee_code="UFOTHER1", employee_name="UF Other Employee",
                plant_location_key="UF Other Plant", company="Ultrafine",
            ),
        ])
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status/company/ROBO/plant/ROBO%20Detail%20Plant")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert "Robo Detail Employee" in html
        assert "UF Other Employee" not in html

    def test_rdc_rejected_as_company_in_company_plant_route(self, client, db, app):
        admin = _make_user("CompanyTabAdmin4", "companytabadmin4@t.com", UserRole.SUPER_ADMIN, db)
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status/company/RDC/plant/Anything")
        assert resp.status_code == 404

    def test_invalid_company_404s(self, client, db, app):
        admin = _make_user("CompanyTabAdmin5", "companytabadmin5@t.com", UserRole.SUPER_ADMIN, db)
        db.session.commit()
        with app.app_context():
            login(client, admin.email)
            resp = client.get("/requests/staffing-status/company/NotAThing/plant/Anything")
        assert resp.status_code == 404


class TestStaffingStatusCompanyScopeGating:
    """
    Regression coverage for the 2026-09-23 fix: the Staffing Status pages
    (and their downloads) predate the company-scope tick-mark feature
    (2026-09-21) and were never updated to respect it — a Business Head or
    HR Manager ticked only for ROBO could still browse full RDC/Ultrafine
    data here, tabs and direct URLs alike. HEAD_HR/DR_BHOON/SUPER_ADMIN stay
    unscoped by design (confirmed with the stakeholder) — not covered here
    since every other test in this file already exercises them.
    """

    def test_robo_only_bh_sees_only_robo_tab(self, client, db, app):
        bh = _make_user("ScopeBh1", "scopebh1@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
        db.session.commit()
        with app.app_context():
            login(client, bh.email)
            resp = client.get("/requests/staffing-status")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert 'id="ctab-btn-robo"' in html
        assert 'id="ctab-btn-rdc"' not in html
        assert 'id="ctab-btn-ultrafine"' not in html

    def test_unscoped_hr_manager_sees_no_tabs(self, client, db, app):
        hrm = _make_user("ScopeHrm1", "scopehrm1@t.com", UserRole.HR_MANAGER, db)  # no companies ticked
        db.session.commit()
        with app.app_context():
            login(client, hrm.email)
            resp = client.get("/requests/staffing-status")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        assert 'id="ctab-btn-rdc"' not in html
        assert 'id="ctab-btn-ultrafine"' not in html
        assert 'id="ctab-btn-robo"' not in html

    def test_robo_only_bh_cannot_reach_ultrafine_plant_detail_by_url(self, client, db, app):
        bh = _make_user("ScopeBh2", "scopebh2@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
        db.session.add(PlantLocation(name="UF Gate Plant", company="Ultrafine", is_active=True))
        db.session.commit()
        with app.app_context():
            login(client, bh.email)
            resp = client.get("/requests/staffing-status/company/Ultrafine/plant/UF%20Gate%20Plant")
        assert resp.status_code == 403

    def test_robo_only_hr_manager_cannot_reach_rdc_plant_detail_by_url(self, client, db, app):
        """Before this fix, HR_MANAGER had ZERO scoping on this route at all
        (only BUSINESS_HEAD's region check existed) — this is the clearest
        regression case."""
        hrm = _make_user("ScopeHrm2", "scopehrm2@t.com", UserRole.HR_MANAGER, db, companies=["ROBO"])
        cluster = ClusterNameMapping(canonical_cluster_name="Gate Cluster")
        db.session.add(cluster)
        db.session.flush()
        db.session.add(PlantDvtMapping(
            plant_location_name="Gate RDC Plant", cluster_id=cluster.id,
            dvt_plant_code="GATE1", match_confidence=MatchConfidence.AUTO_EXACT,
        ))
        db.session.commit()
        with app.app_context():
            login(client, hrm.email)
            resp = client.get("/requests/staffing-status/plant/Gate%20RDC%20Plant")
        assert resp.status_code == 403

    def test_robo_only_hr_manager_cannot_reach_rdc_wide_employee_directory_data(self, client, db, app):
        hrm = _make_user("ScopeHrm3", "scopehrm3@t.com", UserRole.HR_MANAGER, db, companies=["ROBO"])
        db.session.commit()
        db.session.add(EmployeeLocationSnapshot(
            computed_at=datetime.utcnow(), source=ExternalDesignationSource.ZINGHR,
            employee_code="GATERDC1", employee_name="Gate RDC Employee",
            plant_location_key="Some RDC Plant",
        ))
        db.session.commit()
        with app.app_context():
            login(client, hrm.email)
            resp = client.get("/requests/staffing-status/employees")
        assert resp.status_code == 200
        assert "Gate RDC Employee" not in resp.get_data(as_text=True)

    def test_robo_only_bh_cannot_download_ultrafine_report(self, client, db, app):
        bh = _make_user("ScopeBh3", "scopebh3@t.com", UserRole.BUSINESS_HEAD, db, companies=["ROBO"])
        db.session.commit()
        with app.app_context():
            login(client, bh.email)
            resp = client.get("/requests/staffing-status/download/Ultrafine")
        assert resp.status_code == 403

    def test_rdc_ticked_hr_manager_still_sees_rdc_tab(self, client, db, app):
        """Confirms the fix is genuinely scoped, not accidentally fail-closed
        for a legitimately-ticked company."""
        hrm = _make_user("ScopeHrm4", "scopehrm4@t.com", UserRole.HR_MANAGER, db, companies=["RDC"])
        db.session.commit()
        with app.app_context():
            login(client, hrm.email)
            resp = client.get("/requests/staffing-status")
        html = resp.get_data(as_text=True)
        assert 'id="ctab-btn-rdc"' in html
        assert 'id="ctab-btn-ultrafine"' not in html
        assert 'id="ctab-btn-robo"' not in html
