import io
import re
from datetime import datetime, timedelta
from flask import make_response, render_template, request, jsonify
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import joinedload, selectinload
from ..extensions import db
from ..models import (
    OnboardingRequest, RequestStatus, UserRole, User,
    FormField, FieldType, PlantLocation, Designation, ApprovalAction,
)
from ..extensions import limiter
from ..utils import role_required, log_audit, fmt_date
from . import exports_bp

_ALLOWED_ROLES = (UserRole.SUPER_ADMIN, UserRole.HEAD_HR, UserRole.HR_MANAGER, UserRole.DR_BHOON)

_STATUS_LABELS = {
    "DRAFT":             "Draft",
    "PENDING_BH":        "Pending Business / Functional Head",
    "PENDING_DR_BHOON":  "Pending Special Approval",
    "PENDING_HR_MANAGER":"Pending HR Manager",
    "PENDING_HEAD_HR":   "Pending Head HR",
    "ACTIVE":            "Approved",
    "REJECTED_BH":       "Rejected by Business / Functional Head",
    "REJECTED_DR_BHOON": "Rejected by Special Approver",
    "REJECTED_HRM":      "Rejected by HR Manager",
    "REJECTED_HEAD_HR":  "Rejected by Head HR",
}


# ── Shared: collect filter params from request.args ───────────────────────────

def _collect_filter_params(args):
    return {
        "statuses_raw":   args.getlist("status"),
        "date_from_s":    args.get("date_from",    "").strip(),
        "date_to_s":      args.get("date_to",      "").strip(),
        "joining_from_s": args.get("joining_from", "").strip(),
        "joining_to_s":   args.get("joining_to",   "").strip(),
        "plant_filter":   args.get("plant",        "").strip(),
        "desig_filter":   args.get("designation",  "").strip(),
        "company_filter": args.get("company",      "").strip(),
        "initiator_id_s": args.get("initiator_id", "").strip(),
        "bh_id_s":        args.get("bh_id",        "").strip(),
        "retry_min_s":    args.get("retry_min",    "").strip(),
        "retry_max_s":    args.get("retry_max",    "").strip(),
    }


# ── Shared: apply filters and return records ──────────────────────────────────

def _apply_filters(params, viewer=None):
    """Build query from filter params dict; return list of matching records."""
    p = params
    # initiator/actions are lazy="select" by default (actions.actor too) —
    # _cell_val()/_actor_for_role() below touch both per record, which was
    # a real N+1 (2+ extra queries per row). Negligible at today's row
    # counts but compounds linearly as onboarding_requests grows, and it's
    # a one-line fix, so applied eagerly rather than left for later.
    q = OnboardingRequest.query.filter_by(is_deleted=False).options(
        joinedload(OnboardingRequest.initiator),
        selectinload(OnboardingRequest.actions).joinedload(ApprovalAction.actor),
    )

    # Company-scope gating (added 2026-09-23 — same gap class fixed across
    # Staffing Status/admin requests_list/view_request: this export predates
    # the company-scope tick-mark feature). HR_MANAGER is the only scoped
    # role in _ALLOWED_ROLES (SUPER_ADMIN/HEAD_HR/DR_BHOON stay unscoped) —
    # without this, an HR Manager ticked only for ROBO could export a full
    # Excel of every company's candidates, including RDC/Ultrafine PII.
    if viewer is not None and viewer.role == UserRole.HR_MANAGER:
        from ..utils import company_scope_ids
        my_companies = company_scope_ids(viewer.id)
        q = q.filter(OnboardingRequest.company_code.in_(my_companies)) if my_companies else q.filter(db.false())

    # Drafts are private to their Reporting Manager everywhere else in the app - never export them.
    q = q.filter(OnboardingRequest.status != RequestStatus.DRAFT)

    # Status
    if p["statuses_raw"]:
        try:
            status_enums = [RequestStatus(s) for s in p["statuses_raw"] if s]
            if status_enums:
                q = q.filter(OnboardingRequest.status.in_(status_enums))
        except ValueError:
            pass
    else:
        q = q.filter_by(status=RequestStatus.ACTIVE)

    # Submission date range
    if p["date_from_s"]:
        try:
            q = q.filter(OnboardingRequest.created_at >=
                         datetime.strptime(p["date_from_s"], "%Y-%m-%d"))
        except ValueError:
            pass
    if p["date_to_s"]:
        try:
            q = q.filter(OnboardingRequest.created_at <
                         datetime.strptime(p["date_to_s"], "%Y-%m-%d") + timedelta(days=1))
        except ValueError:
            pass

    # Plant / designation / company
    if p["plant_filter"]:
        q = q.filter_by(plant_location=p["plant_filter"])
    if p["desig_filter"]:
        q = q.filter_by(designation=p["desig_filter"])
    if p["company_filter"]:
        q = q.filter_by(company_code=p["company_filter"])

    # Initiator / BH
    if p["initiator_id_s"]:
        try:
            q = q.filter_by(initiated_by=int(p["initiator_id_s"]))
        except ValueError:
            pass
    elif p["bh_id_s"]:
        try:
            bh_int = int(p["bh_id_s"])
            # NOT business_head_id (legacy/unwritten since the 2026-09-04
            # region-routing redesign — see that column's docstring in
            # models.py). Real BH eligibility is per-request, since it
            # depends on both the initiator's regions and the request's own
            # company_code (bh_ids_for_initiator's signature). Re-derive it
            # per company via the same canonical routing function so this
            # filter can never drift from actual approval-routing behavior.
            from ..utils import bh_ids_for_initiator
            from ..models import COMPANY_CHOICES
            # HR Managers can also initiate requests (2026-09-29) — must be
            # included here too, or a BH-filtered export would silently
            # drop every HR-Manager-initiated request routed to that BH.
            all_initiators = User.query.filter(
                User.role.in_([UserRole.INITIATOR, UserRole.HR_MANAGER])
            ).all()
            clauses = []
            for company in COMPANY_CHOICES:
                ids = [u.id for u in all_initiators if bh_int in bh_ids_for_initiator(u, company)]
                if ids:
                    clauses.append(db.and_(
                        OnboardingRequest.company_code == company,
                        OnboardingRequest.initiated_by.in_(ids),
                    ))
            q = q.filter(db.or_(*clauses)) if clauses else q.filter(OnboardingRequest.id == -1)
        except ValueError:
            pass

    # Retry count range
    if p["retry_min_s"]:
        try:
            q = q.filter(OnboardingRequest.retry_count >= int(p["retry_min_s"]))
        except ValueError:
            pass
    if p["retry_max_s"]:
        try:
            q = q.filter(OnboardingRequest.retry_count <= int(p["retry_max_s"]))
        except ValueError:
            pass

    records = q.order_by(OnboardingRequest.created_at.desc()).all()

    # Python-side joining date filter (stored in JSON form_data)
    if p["joining_from_s"] or p["joining_to_s"]:
        try:
            jf = datetime.strptime(p["joining_from_s"], "%Y-%m-%d").date() if p["joining_from_s"] else None
            jt = datetime.strptime(p["joining_to_s"],   "%Y-%m-%d").date() if p["joining_to_s"]   else None
            filtered = []
            for r in records:
                jd_str = r.form_data.get("contract_from") or r.form_data.get("joining_date", "")   # the form field is contract_from
                if not jd_str:
                    continue
                try:
                    jd = datetime.strptime(jd_str[:10], "%Y-%m-%d").date()
                    if jf and jd < jf:
                        continue
                    if jt and jd > jt:
                        continue
                    filtered.append(r)
                except Exception:
                    continue
            records = filtered
        except Exception:
            pass

    return records


# ── Shared: report title ──────────────────────────────────────────────────────

def _report_title(params):
    if params["statuses_raw"]:
        status_str = " + ".join(_STATUS_LABELS.get(s, s) for s in params["statuses_raw"])
    else:
        status_str = "Approved"
    return f"{status_str} — Onboarding Report"


# ── Shared: human-readable filter summary chips ───────────────────────────────

def _filter_chips(params):
    chips = []
    p = params

    if p["statuses_raw"]:
        labels = [_STATUS_LABELS.get(s, s) for s in p["statuses_raw"]]
        chips.append({"label": "Status", "value": ", ".join(labels)})
    else:
        chips.append({"label": "Status", "value": "Approved (default)"})

    if p["date_from_s"] or p["date_to_s"]:
        val = f"{fmt_date(p['date_from_s']) or '…'} → {fmt_date(p['date_to_s']) or '…'}"
        chips.append({"label": "Submitted", "value": val})

    if p["joining_from_s"] or p["joining_to_s"]:
        val = f"{fmt_date(p['joining_from_s']) or '…'} → {fmt_date(p['joining_to_s']) or '…'}"
        chips.append({"label": "Joining date", "value": val})

    if p["plant_filter"]:
        chips.append({"label": "Plant", "value": p["plant_filter"]})

    if p["desig_filter"]:
        chips.append({"label": "Designation", "value": p["desig_filter"]})

    if p["company_filter"]:
        chips.append({"label": "Company", "value": p["company_filter"]})

    if p["initiator_id_s"]:
        u = db.session.get(User, int(p["initiator_id_s"])) if p["initiator_id_s"].isdigit() else None
        chips.append({"label": "Reporting Manager", "value": u.name if u else p["initiator_id_s"]})

    if p["bh_id_s"]:
        u = db.session.get(User, int(p["bh_id_s"])) if p["bh_id_s"].isdigit() else None
        chips.append({"label": "Business / Functional Head", "value": u.name if u else p["bh_id_s"]})

    if p["retry_min_s"] or p["retry_max_s"]:
        val = f"{p['retry_min_s'] or '0'} – {p['retry_max_s'] or '∞'}"
        chips.append({"label": "Retries", "value": val})

    return chips


# ── Shared: get ordered form field columns (no FILE fields) ───────────────────

def _get_form_fields():
    return FormField.query.filter(
        FormField.is_active == True,
        FormField.is_deleted == False,
        FormField.field_type != FieldType.FILE,
    ).order_by(FormField.step, FormField.sort_order).all()


# ── Shared: format a single record cell value ─────────────────────────────────

def _actor_for_role(req, role_value):
    """Return 'Name (EMP_CODE)' for the first APPROVED action by a given role."""
    from ..models import ApprovalActionType
    for act in req.actions:
        if (act.actor and act.actor.role.value == role_value
                and act.action == ApprovalActionType.APPROVED):
            code = act.actor.employee_code or "—"
            return f"{act.actor.name} ({code})"
    return ""


def _plant_display_name(plant_location_name, cache):
    """Raw plant_location_name -> proper ERP-Tracker display name (see
    PlantDvtMapping.display_name), cached across one export run so a
    multi-hundred-row report doesn't re-query per row."""
    if not plant_location_name:
        return plant_location_name
    if plant_location_name in cache:
        return cache[plant_location_name]
    from ..models import PlantDvtMapping
    row = PlantDvtMapping.query.filter_by(plant_location_name=plant_location_name, is_deleted=False).first()
    result = row.display_name if row else plant_location_name
    cache[plant_location_name] = result
    return result


def _cell_val(req, key, plant_name_cache=None):
    if   key == "_id":          return req.id
    elif key == "_status":      return _STATUS_LABELS.get(req.status.value, req.status.value)
    elif key == "_initiator":   return req.initiator.name if req.initiator else ""
    elif key == "_init_code":
        return (req.initiator.employee_code or "—") if req.initiator else ""
    elif key == "_bh_approval": return _actor_for_role(req, "BUSINESS_HEAD")
    elif key == "_hrm_approval":return _actor_for_role(req, "HR_MANAGER")
    elif key == "_hhr_approval":return _actor_for_role(req, "HEAD_HR")
    elif key == "_sa_approval": return _actor_for_role(req, "DR_BHOON")
    elif key == "_created":
        return req.created_at.strftime("%d/%m/%Y") if req.created_at else ""
    elif key == "_updated":
        return req.updated_at.strftime("%d/%m/%Y") if req.updated_at else ""
    elif key == "_retry":       return req.retry_count
    else:
        fd  = req.form_data
        raw = fd.get(key, "")
        if raw == "Other":
            other = fd.get(f"{key}_other", "")
            return f"Other: {other}" if other else "Other"
        if key == "plant_location" and raw:
            return _plant_display_name(raw, plant_name_cache if plant_name_cache is not None else {})
        if isinstance(raw, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw.strip()):
            return fmt_date(raw)         # date fields (contract_from, date_of_birth, ...) are stored ISO
        return raw


# ── Excel builder ──────────────────────────────────────────────────────────────

def _build_excel(records, report_title="Employee Onboarding Report"):
    form_fields = _get_form_fields()

    meta_cols = [
        ("Request ID",        "_id",          12),
        ("Status",            "_status",      22),
        ("Initiated By",      "_initiator",   22),
        ("RM Emp Code","_init_code",   18),
        ("Business / Functional Head Approval", "_bh_approval", 28),
        ("HR Manager Approval",      "_hrm_approval",28),
        ("Head HR Approval",         "_hhr_approval",28),
        ("Special Approver Approval","_sa_approval", 28),
        ("Submitted On",      "_created",     18),
        ("Last Updated",      "_updated",     18),
        ("Retry Count",       "_retry",       12),
    ]
    form_cols = [(f.field_label, f.field_key, max(len(f.field_label) + 4, 18)) for f in form_fields]
    all_cols = meta_cols + form_cols

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    hdr_fill  = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    hdr_font  = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dat_font  = Font(name="Calibri", size=10)
    dat_align = Alignment(vertical="center", wrap_text=False)
    thin      = Side(style="thin", color="CCCCCC")
    bdr       = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt_fill  = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    # Title row
    ws.merge_cells(f"A1:{get_column_letter(len(all_cols))}1")
    tc = ws["A1"]
    tc.value = f"{report_title}  |  Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    tc.font  = Font(name="Calibri", size=13, bold=True, color="1F3864")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Header row
    for ci, (label, _, width) in enumerate(all_cols, 1):
        c = ws.cell(row=2, column=ci, value=label)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = hdr_align; c.border = bdr
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"

    plant_name_cache = {}
    for ri, req in enumerate(records, 3):
        for ci, (_, key, _) in enumerate(all_cols, 1):
            val = _cell_val(req, key, plant_name_cache)
            c = ws.cell(row=ri, column=ci, value=val)
            if isinstance(val, str) and val[:1] in ("=", "+", "-", "@"):
                c.data_type = "s"   # user text must never be evaluated as a spreadsheet formula
            c.font = dat_font; c.alignment = dat_align; c.border = bdr
            if ri % 2 == 0:
                c.fill = alt_fill
        ws.row_dimensions[ri].height = 18

    # Footer
    fr = len(records) + 3
    ws.merge_cells(f"A{fr}:{get_column_letter(len(all_cols))}{fr}")
    fc = ws[f"A{fr}"]
    fc.value = f"Total records: {len(records)}"
    fc.font = Font(name="Calibri", size=10, bold=True, color="555555")
    fc.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Filter page ────────────────────────────────────────────────────────────────

@exports_bp.route("/active-employees")
@login_required
@role_required(*_ALLOWED_ROLES)
def active_employees():
    plants = PlantLocation.query.filter_by(
        is_deleted=False, is_active=True
    ).order_by(PlantLocation.name).all()

    # Designations are per-company now — the same name can appear once per
    # company, so de-dupe by name for this company-agnostic filter dropdown.
    _seen_desig = set()
    designations = []
    for _d in Designation.query.filter_by(is_deleted=False, is_active=True).order_by(Designation.name).all():
        if _d.name not in _seen_desig:
            _seen_desig.add(_d.name)
            designations.append(_d)

    companies = [r[0] for r in db.session.execute(db.text(
        "SELECT DISTINCT company_code FROM onboarding_requests "
        "WHERE is_deleted=0 AND company_code IS NOT NULL AND company_code != '' "
        "ORDER BY company_code"
    )).fetchall()]
    if current_user.role == UserRole.HR_MANAGER:
        from ..utils import company_scope_ids
        my_companies = company_scope_ids(current_user.id)
        companies = [c for c in companies if c in my_companies]

    # HR Managers can also initiate requests (2026-09-29) — include them so
    # a request they submitted can be found by initiator in this filter too.
    initiators = User.query.filter(
        User.role.in_([UserRole.INITIATOR, UserRole.HR_MANAGER]), User.is_active == True  # noqa: E712
    ).order_by(User.name).all()

    bhs = User.query.filter_by(
        role=UserRole.BUSINESS_HEAD, is_active=True
    ).order_by(User.name).all()

    return render_template(
        "exports/filter.html",
        statuses=list(RequestStatus),
        status_labels=_STATUS_LABELS,
        plants=plants,
        designations=designations,
        companies=companies,
        initiators=initiators,
        business_heads=bhs,
    )


# ── Preview endpoint (AJAX) ────────────────────────────────────────────────────

_PREVIEW_ROWS = 10   # rows shown in the preview table

@exports_bp.route("/active-employees/preview")
@login_required
@role_required(*_ALLOWED_ROLES)
@limiter.limit("60 per minute")
def preview_excel():
    """Return JSON preview: columns, first N rows, total count, filter chips."""
    params  = _collect_filter_params(request.args)
    records = _apply_filters(params, viewer=current_user)
    title   = _report_title(params)
    chips   = _filter_chips(params)

    form_fields = _get_form_fields()
    meta_cols = [
        ("Request ID",        "_id"),
        ("Status",            "_status"),
        ("Initiated By",      "_initiator"),
        ("RM Emp Code","_init_code"),
        ("Business / Functional Head Approval", "_bh_approval"),
        ("HR Manager Approval",      "_hrm_approval"),
        ("Head HR Approval",         "_hhr_approval"),
        ("Special Approver Approval","_sa_approval"),
        ("Submitted On",      "_created"),
        ("Last Updated",      "_updated"),
        ("Retry Count",       "_retry"),
    ]
    form_cols = [(f.field_label, f.field_key) for f in form_fields]
    all_cols  = meta_cols + form_cols

    columns   = [label for label, _ in all_cols]
    preview   = records[:_PREVIEW_ROWS]
    plant_name_cache = {}
    rows      = [
        [str(_cell_val(r, key, plant_name_cache)) for _, key in all_cols]
        for r in preview
    ]

    return jsonify({
        "ok":           True,
        "total":        len(records),
        "preview_rows": _PREVIEW_ROWS,
        "title":        title,
        "columns":      columns,
        "rows":         rows,
        "chips":        chips,
    })


# ── Download endpoint ──────────────────────────────────────────────────────────

@exports_bp.route("/active-employees/download")
@login_required
@role_required(*_ALLOWED_ROLES)
def download_excel():
    params  = _collect_filter_params(request.args)
    records = _apply_filters(params, viewer=current_user)
    title   = _report_title(params)

    buf  = _build_excel(records, title)
    data = buf.getvalue()

    slug     = "_".join(params["statuses_raw"]) if params["statuses_raw"] else "ACTIVE"
    filename = f"onboarding_{slug}_{datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"

    _filters = {k: v for k, v in params.items() if v}
    log_audit("EXPORT", "EXPORT_DOWNLOADED",
              detail={"record_count": len(records),
                      "filename":     filename,
                      "report_title": title,
                      "filters":      _filters})
    db.session.commit()

    resp = make_response(data)
    resp.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    resp.headers["Content-Length"]       = len(data)
    resp.headers["Cache-Control"]        = "no-cache, no-store, must-revalidate"
    resp.headers["Pragma"]               = "no-cache"
    resp.headers["Expires"]              = "0"
    return resp


# ── Action times: every recorded step with its time, for response-time analysis ─
# (2026-10-14) One row per RequestEvent - submitted / approved / rejected / resubmitted / Truein push ... -
# with the exact time, who did it, and how long the request had waited for it. Second sheet: average
# response time per role (the number the stakeholders asked for). Same audience as the other export.

_IST = timedelta(hours=5, minutes=30)


def _ist_text(dt):
    return (dt + _IST).strftime("%d/%m/%Y %H:%M:%S") if dt else ""


def _parse_iso_day(value):
    try:
        return datetime.strptime((value or "").strip(), "%Y-%m-%d")
    except ValueError:
        return None


def _minutes(seconds):
    return round(seconds / 60.0, 2) if seconds is not None else None


@exports_bp.route("/action-times/download")
@login_required
@role_required(*_ALLOWED_ROLES)
def download_action_times():
    from statistics import median
    from ..models import RequestEvent, ROLE_LABELS, STATUS_LABELS
    from ..utils import EVENT_LABELS, format_duration

    d_from = _parse_iso_day(request.args.get("from"))
    d_to = _parse_iso_day(request.args.get("to"))
    company = (request.args.get("company") or "").strip()

    q = (db.session.query(RequestEvent, OnboardingRequest)
         .join(OnboardingRequest, OnboardingRequest.id == RequestEvent.request_id))
    if d_from:                                    # the chosen days are IST days; events are stored in UTC
        q = q.filter(RequestEvent.occurred_at >= d_from - _IST)
    if d_to:
        q = q.filter(RequestEvent.occurred_at < d_to + timedelta(days=1) - _IST)
    if current_user.role == UserRole.HR_MANAGER:
        from ..utils import company_scope_ids
        q = q.filter(OnboardingRequest.company_code.in_(company_scope_ids(current_user.id) or [""]))
    if company:
        q = q.filter(OnboardingRequest.company_code == company)
    rows = q.order_by(RequestEvent.request_id, RequestEvent.occurred_at, RequestEvent.id).all()

    def stage(v):
        return STATUS_LABELS.get(v, v or "")

    wb = Workbook()
    hdr_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dat_font = Font(name="Calibri", size=10)
    thin = Side(style="thin", color="CCCCCC")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

    def write_sheet(ws, title, columns, data):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        c = ws.cell(row=1, column=1, value=f"{title}  |  Generated: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        c.font = Font(name="Calibri", size=13, bold=True, color="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        for ci, (label, width) in enumerate(columns, 1):
            h = ws.cell(row=2, column=ci, value=label)
            h.font, h.fill, h.alignment, h.border = hdr_font, hdr_fill, hdr_align, bdr
            ws.column_dimensions[get_column_letter(ci)].width = width
        ws.row_dimensions[2].height = 32
        ws.freeze_panes = "A3"
        for ri, vals in enumerate(data, 3):
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row=ri, column=ci, value=v)
                if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
                    cell.data_type = "s"          # user text is never evaluated as a formula
                cell.font, cell.border = dat_font, bdr

    # Sheet 1: every event
    ws = wb.active
    ws.title = "Action log"
    cols = [("Request ID", 11), ("Candidate", 24), ("Company", 11), ("Plant", 22), ("Designation", 22),
            ("Attempt", 9), ("Action", 30), ("Done by", 22), ("Role", 24), ("From stage", 30), ("To stage", 30),
            ("Date & time (IST)", 20), ("Request had been waiting since (IST)", 22), ("Time taken", 14),
            ("Time taken (minutes)", 14), ("Total approval time so far (minutes)", 18), ("How", 12), ("Remark", 50)]
    data = []
    for ev, rq in rows:
        data.append([
            rq.id, rq.candidate_name or "", rq.company_code or "", rq.plant_location or "", rq.designation or "",
            ev.attempt_no, EVENT_LABELS.get(ev.event_type, ev.event_type), ev.actor_name or "System",
            ROLE_LABELS.get(ev.actor_role, ev.actor_role or ""), stage(ev.from_status), stage(ev.to_status),
            _ist_text(ev.occurred_at), _ist_text(ev.waiting_since),
            format_duration(ev.duration_seconds) if ev.duration_seconds is not None else "",
            _minutes(ev.duration_seconds), _minutes(ev.cycle_seconds),
            {"email_link": "E-mail link", "system": "Automatic", "manual": "Manual"}.get(ev.via, "App"),
            ev.remark or ""])
    write_sheet(ws, "Request action log - every step with its time", cols, data)

    # Sheet 2: response time by role (decisions only: approved / rejected)
    groups = {}
    for ev, rq in rows:
        if ev.event_type in ("APPROVED", "REJECTED") and ev.duration_seconds is not None:
            for co in (rq.company_code or "", "All companies"):
                groups.setdefault((ev.actor_role or "", co), []).append((ev.event_type, ev.duration_seconds))
    stat_rows = []
    for (role, co), items in sorted(groups.items(), key=lambda kv: (kv[0][1] != "All companies", kv[0][1], kv[0][0])):
        secs = [s for _, s in items]
        stat_rows.append([ROLE_LABELS.get(role, role), co, len(items),
                          sum(1 for t, _ in items if t == "APPROVED"), sum(1 for t, _ in items if t == "REJECTED"),
                          format_duration(sum(secs) / len(secs)), _minutes(sum(secs) / len(secs)),
                          format_duration(median(secs)), _minutes(median(secs)),
                          format_duration(min(secs)), format_duration(max(secs))])
    ws2 = wb.create_sheet("Response time by role")
    write_sheet(ws2, "How long each role takes to respond (from the request reaching them to their decision)",
                [("Role", 26), ("Company", 14), ("Decisions", 11), ("Approved", 10), ("Rejected", 10),
                 ("Average", 16), ("Average (minutes)", 14), ("Median", 16), ("Median (minutes)", 14),
                 ("Fastest", 16), ("Slowest", 16)], stat_rows)

    buf = io.BytesIO()
    wb.save(buf)
    data_bytes = buf.getvalue()
    filename = f"action_times_{datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
    log_audit("EXPORT", "EXPORT_DOWNLOADED",
              detail={"report_title": "Action times", "record_count": len(rows), "filename": filename,
                      "filters": {"from": request.args.get("from"), "to": request.args.get("to"), "company": company}})
    db.session.commit()

    resp = make_response(data_bytes)
    resp.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    resp.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    resp.headers["Content-Length"] = len(data_bytes)
    resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return resp
